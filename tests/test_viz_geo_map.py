"""F11 地域分布：静态地图 PNG + Word 附录装配的回归保护。

两条底线：
1. 有地域数据 → 出表出图；无地域数据/无底图 → 一个字都不写（不静默造图）。
2. geo_map 缺省时，Word 输出与既有逐字一致（老调用方零感知）。
"""

import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from parser import parse_report  # noqa: E402
from viz_network import _generate_geo_png, _geo_polygons, _locate_geojson, generate_diagram  # noqa: E402


GEO_SAMPLE = {
    "regions": [
        {"region_name": "广东省", "independent_sources": 9, "sources": 14, "share": 0.45,
         "cities": ["深圳市", "广州市"]},
        {"region_name": "北京市", "independent_sources": 5, "sources": 7, "share": 0.25,
         "cities": ["北京市"]},
    ],
    "coverage": 0.7,
}


class GeoPolygonParsingTests(unittest.TestCase):
    def test_polygon_and_multipolygon_rings_are_flattened(self):
        polygon = {"type": "Polygon", "coordinates": [[[110.0, 20.0], [111.0, 20.0], [111.0, 21.0]]]}
        multi = {"type": "MultiPolygon", "coordinates": [
            [[[110.0, 20.0], [111.0, 20.0], [111.0, 21.0]]],
            [[[112.0, 22.0], [113.0, 22.0], [113.0, 23.0]]],
        ]}
        self.assertEqual(len(_geo_polygons(polygon)), 1)
        self.assertEqual(len(_geo_polygons(multi)), 2)

    def test_degenerate_ring_is_dropped(self):
        # 3 点以下无法成面，必须丢弃而不是抛异常
        tiny = {"type": "Polygon", "coordinates": [[[110.0, 20.0], [111.0, 20.0]]]}
        self.assertEqual(_geo_polygons(tiny), [])
        self.assertEqual(_geo_polygons({}), [])
        self.assertEqual(_geo_polygons(None), [])


class GeolocatedPngTests(unittest.TestCase):
    def test_basemap_is_locatable(self):
        # 底图是出图前提；找不到就必须返回 None 而不是画张空图
        self.assertTrue(_locate_geojson(), "未能定位 GeoJSON 底图")

    def test_renders_when_regions_present(self):
        with tempfile.TemporaryDirectory() as folder:
            out = os.path.join(folder, "geo.png")
            self.assertIsNotNone(generate_diagram({"viz": "geo", **GEO_SAMPLE}, out))
            self.assertGreater(os.path.getsize(out), 20_000)
            with open(out, "rb") as handle:
                self.assertEqual(handle.read(8), b"\x89PNG\r\n\x1a\n")

    def test_returns_none_without_regions(self):
        with tempfile.TemporaryDirectory() as folder:
            out = os.path.join(folder, "empty.png")
            self.assertIsNone(generate_diagram({"viz": "geo", "regions": []}, out))
            self.assertFalse(os.path.exists(out))

    def test_returns_none_when_basemap_missing(self):
        import viz_network
        original = viz_network._locate_geojson
        viz_network._locate_geojson = lambda: None
        try:
            with tempfile.TemporaryDirectory() as folder:
                out = os.path.join(folder, "nobase.png")
                self.assertIsNone(_generate_geo_png(GEO_SAMPLE, out))
        finally:
            viz_network._locate_geojson = original


class WordGeoAppendixTests(unittest.TestCase):
    MD = (
        "# 测试报告\n\n"
        "## 情况概述\n\n概述正文\n\n"
        "## 分析框架\n\n框架正文\n\n"
        "## 三元结构分析正文\n\n分析正文\n\n"
        "## 结论\n\n结论正文\n\n"
        "## 附录\n\n[来源一](https://example.com/1)\n"
    )

    def _headings(self, docx_path):
        return [p.text for p in Document(docx_path).paragraphs if p.style.name == "Heading 1"]

    def _table_text(self, docx_path):
        return "\n".join(
            cell.text for table in Document(docx_path).tables
            for row in table.rows for cell in row.cells
        )

    def test_geo_map_renders_heading_table_and_picture(self):
        from docx_renderer import render_docx
        report = parse_report(self.MD)
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "geo.docx"
            render_docx(report, str(out), output_folder=folder, geo_map=dict(GEO_SAMPLE))
            doc = Document(out)
            self.assertIn("附表 1：来源地域分布", [p.text for p in doc.paragraphs if p.style.name == "Heading 1"])
            self.assertEqual(len(doc.inline_shapes), 1, "地域分布地图未插入")
            table_text = self._table_text(out)
            self.assertIn("广东省", table_text)
            self.assertIn("独立源数", table_text)  # 表头
            self.assertIn("深圳市", table_text)  # 涉及地市
            self.assertTrue(os.path.exists(os.path.join(folder, "地域分布.png")))

    def test_appendix_number_continues_after_data_tables(self):
        from docx_renderer import render_docx
        report = parse_report(self.MD)
        # 首张是空表（会被跳过）——地域分布的编号必须按「实际渲染数」接续，不能按入参长度
        tables = [
            {"name": "空表", "columns": [], "rows": []},
            {"name": "关键数据表", "columns": ["项", "值"], "rows": [["甲", "1"]]},
        ]
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "both.docx"
            render_docx(report, str(out), output_folder=folder, data_tables=tables, geo_map=dict(GEO_SAMPLE))
            headings = self._headings(out)
            self.assertIn("附表 1：关键数据表", headings)
            self.assertNotIn("附表 1：空表", headings)
            self.assertIn("附表 2：来源地域分布", headings)

    def test_output_unchanged_without_geo_map(self):
        from docx_renderer import render_docx
        report = parse_report(self.MD)
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "plain.docx"
            render_docx(report, str(out), output_folder=folder)
            headings = self._headings(out)
            self.assertNotIn("附表 1：来源地域分布", headings)
            self.assertEqual(len(Document(out).inline_shapes), 0)

    def test_empty_regions_writes_nothing(self):
        from docx_renderer import render_docx
        report = parse_report(self.MD)
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder) / "emptygeo.docx"
            render_docx(report, str(out), output_folder=folder, geo_map={"regions": [], "coverage": 0.0})
            headings = self._headings(out)
            self.assertNotIn("附表 1：来源地域分布", headings)
            self.assertEqual(len(Document(out).inline_shapes), 0)


if __name__ == "__main__":
    unittest.main()
