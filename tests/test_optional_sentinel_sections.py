"""可选哨兵章节回归（F15 账本 1.3 配套）：

- 「叙事份额分析」(opinion_narrative_share) 与「政策条款拆解」(policy_clause_breakdown)
  是**可选节**：不进 SENTINEL_SECTIONS / 路由判定，正文不写 → Word 输出与既有逐字一致；
  正文写了 → 按模式槽位插入并参与统一章节编号。
- Word 附表「关键数据表」（data_tables）为可选通道：None 时不渲染附表。
"""
import re

from docx import Document

from engine import CaseAnalysisEngine


_OPINION_BODY_BASE = """# 叙事份额回归测试

## 情况概述

这是一场舆论事件的范围交代与结论前置。

## 事件与时间线

1. 某事件在社交平台发酵。（来源：测试来源）

## 利益主体与沉默方

显性发声者为当事方，沉默方为未被代表的使用者。

## 叙事竞争矩阵

A 叙事与 B 叙事争夺定义权。

## 三元生命维度

这场舆论的定义权锚点、信任链与防御协议如下。

## 逆反性质与层级

当前落在群体层级的建设性区间。

## 演化曲线与系统回应

处于扩散节点，系统回应趋向修复回路。

## 核心冲突点

1. 【当事方】对【使用者】在【解释权】上的张力：叙事归属冲突。

## 结论

### 汇流段

当事方与使用者之争本质是定义权之争。

### 核心判断

舆论的走向取决于系统回应的质量。

### 博弈终局预判

1. 回应透明后热度回落。

## 行动建议

- 对当事方：主动披露事实，保留回应记录。

## 附录

[测试来源](https://example.com/source)
"""

_NARRATIVE_SECTION = """## 叙事份额分析

A 叙事约 45%（口径：样本 20 条来源中 9 条持此叙事）。

"""

_POLICY_BODY_BASE = """# 政策条款回归测试

## 情况概述

这是政策分析的范围交代与结论前置。

## 独立事实摘要

1. 某市印发企业奖补措施。（来源：测试来源）

## 分析框架说明

这不是普惠补贴，而是定向竞争资源的配置。

## 政策对象图谱

### 基本信息

发文主体为市政府。

### 受影响群体（四分法）

直接受益者为中小企业。

## 政策权重与空间分析

### 权重层级判定

市级政策，权重中等。

### 操作空间评估

申领窗口与材料准备存在操作空间。

## 核心冲突点

1. 【企业】对【审批部门】在【奖补资金】上的张力：材料门槛与合规成本冲突。

## 三元结构分析正文

### 第一节：奖补的分配逻辑

奖补资金不是普惠发放，而是定向激励。

## 结论与推导

### 汇流段

企业与审批部门围绕材料门槛展开博弈。

### 博弈终局预判

1. 材料标准化后申领成本下降。

## 行动建议

- 对企业：提前准备审计材料，跟踪申领窗口。

## 附录/数据溯源

[测试来源](https://example.com/source)
"""

_CLAUSE_SECTION = """## 政策条款拆解

### 条款定位

第二条：符合条件的企业可申领最高 50 万元奖补。

### 影响对象与利益变化

中小企业受益，个体工商户间接受益。

### 生效与地域层级

市级，2026-09-01 起施行。

"""


def _h1_texts(docx_path: str) -> list[str]:
    doc = Document(docx_path)
    return [
        p.text
        for p in doc.paragraphs
        if p.style is not None and p.style.name == "Heading 1"
    ]


def _insert_after(body: str, anchor: str, insert: str) -> str:
    marker = f"## {anchor}"
    idx = body.index(marker)
    line_end = body.index("\n", idx) + 1
    return body[:line_end] + "\n" + insert + body[line_end:]


def test_opinion_without_optional_section_has_no_share_heading(tmp_path):
    """不写可选节：输出无「叙事份额分析」，其余章节编号不变（零回归）。"""
    result = CaseAnalysisEngine().export_from_text(
        "叙事份额回归测试", _OPINION_BODY_BASE,
        output_dir=str(tmp_path), slug="opinion-base",
    )
    h1 = _h1_texts(result["word"])
    assert all("叙事份额分析" not in t for t in h1)
    # 相邻章节编号连续：叙事竞争矩阵 → 三元生命维度
    idx = next(i for i, t in enumerate(h1) if "叙事竞争矩阵" in t)
    assert "三元生命维度" in h1[idx + 1]


def test_opinion_with_optional_section_inserted_between(tmp_path):
    """写可选节：插入叙事竞争矩阵与三元生命维度之间，编号连续无重复。"""
    body = _insert_after(_OPINION_BODY_BASE, "叙事竞争矩阵", _NARRATIVE_SECTION)
    result = CaseAnalysisEngine().export_from_text(
        "叙事份额回归测试", body,
        output_dir=str(tmp_path), slug="opinion-share",
    )
    h1 = _h1_texts(result["word"])
    share_idx = next(i for i, t in enumerate(h1) if "叙事份额分析" in t)
    assert "叙事竞争矩阵" in h1[share_idx - 1]
    assert "三元生命维度" in h1[share_idx + 1]
    # 渲染器统一编号必须为连续中文序号（一、二、三……到附录），无跳号无重复；
    # 唯一无编号的 Heading 1 是目录页「目  录」
    numbers = [re.match(r"^([一二三四五六七八九十]+)、", t) for t in h1]
    seq = [m.group(1) for m in numbers if m]
    unnumbered = [t for t, m in zip(h1, numbers) if not m]
    assert unnumbered == ["目  录"]
    assert _to_int(seq[0]) == 1
    assert all(_to_int(b) == _to_int(a) + 1 for a, b in zip(seq, seq[1:]))


def _to_int(cn: str) -> int:
    """极简中文序号转 int（测试用，覆盖到 二十）。"""
    table = {c: i for i, c in enumerate("零一二三四五六七八九")}
    if cn == "十":
        return 10
    if "十" in cn:
        tens, _, ones = cn.partition("十")
        return table.get(tens, 1) * 10 + table.get(ones, 0)
    return table[cn]


def test_policy_without_and_with_clause_section(tmp_path):
    """政策条款拆解可选节：不写不出现；写了插入对象图谱与权重空间之间。"""
    engine = CaseAnalysisEngine()
    base = engine.export_from_text(
        "政策条款回归测试", _POLICY_BODY_BASE,
        output_dir=str(tmp_path), slug="policy-base",
    )
    h1 = _h1_texts(base["word"])
    assert all("政策条款拆解" not in t for t in h1)
    idx = next(i for i, t in enumerate(h1) if "政策对象图谱" in t)
    assert "政策权重与空间分析" in h1[idx + 1]

    body = _insert_after(_POLICY_BODY_BASE, "政策对象图谱", _CLAUSE_SECTION)
    with_clause = CaseAnalysisEngine().export_from_text(
        "政策条款回归测试", body,
        output_dir=str(tmp_path), slug="policy-clause",
    )
    h1 = _h1_texts(with_clause["word"])
    clause_idx = next(i for i, t in enumerate(h1) if "政策条款拆解" in t)
    assert "政策对象图谱" in h1[clause_idx - 1]
    assert "政策权重与空间分析" in h1[clause_idx + 1]


def test_data_tables_appendix_rendered_only_when_provided(tmp_path):
    """Word 附表通道：data_tables=None 无附表；传入时按序渲染「附表 N」。"""
    engine = CaseAnalysisEngine()
    tables = [
        {
            "name": "来源证据表",
            "columns": ["编号", "标题", "链接"],
            "rows": [["s1", "官方公告", "https://example.com/a"]],
        },
        {
            "name": "主体清单表",
            "columns": ["编号", "主体", "角色"],
            "rows": [["n1", "企业", "subject"]],
        },
    ]
    with_tables = engine.export_from_text(
        "附表通道测试", _OPINION_BODY_BASE,
        output_dir=str(tmp_path), slug="with-tables",
        data_tables=tables,
    )
    doc = Document(with_tables["word"])
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "附表 1：来源证据表" in text
    assert "附表 2：主体清单表" in text
    # 表格内容落进 Word 表格对象
    assert any(cell.text == "官方公告" for table in doc.tables for row in table.rows for cell in row.cells)

    without_tables = CaseAnalysisEngine().export_from_text(
        "附表通道测试", _OPINION_BODY_BASE,
        output_dir=str(tmp_path), slug="no-tables",
    )
    doc2 = Document(without_tables["word"])
    text2 = "\n".join(p.text for p in doc2.paragraphs)
    assert "附表" not in text2
